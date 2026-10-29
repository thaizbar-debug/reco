# Genera los archivos de prueba de importación de documentos (tests/import-documents.spec.js).
# Uso: python3 make_fixtures.py <salida> <carpeta temporal>  (requiere Pillow, python-docx y python-pptx)
# El PDF se imprime después con make_pdf.js a partir de <carpeta temporal>/doc.html.
import io, base64, sys
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches, Pt

OUT = sys.argv[1]
SCR = sys.argv[2]

def photo(color, label, size=(480, 360)):
    im = Image.new('RGB', size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([20, 20, size[0]-20, size[1]-20], outline=(255, 255, 255), width=6)
    d.text((40, 40), label, fill=(255, 255, 255))
    b = io.BytesIO(); im.save(b, 'JPEG', quality=80); return b.getvalue()

def logo():
    im = Image.new('RGB', (48, 48), (200, 30, 30))
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()

def scan_page():
    # Página "escaneada": el texto es parte de la imagen, no hay capa de texto.
    im = Image.new('RGB', (900, 500), (250, 250, 245))
    d = ImageDraw.Draw(im)
    try:
        f = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 34)
    except Exception:
        f = ImageFont.load_default()
    for i, line in enumerate(['TERRENO EN VENTA - LURIN', 'Lote de 500 m2', 'Precio: US$ 90,000', 'Zonificacion residencial']):
        d.text((40, 40 + i * 70), line, fill=(20, 20, 20), font=f)
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()

sala, cocina, fachada, playa = photo((60, 90, 160), 'SALA'), photo((160, 90, 60), 'COCINA'), photo((60, 140, 80), 'FACHADA'), photo((40, 120, 170), 'PLAYA')
open(f'{SCR}/sala.jpg', 'wb').write(sala); open(f'{SCR}/cocina.jpg', 'wb').write(cocina)
open(f'{SCR}/fachada.jpg', 'wb').write(fachada); open(f'{SCR}/logo.png', 'wb').write(logo())
open(f'{SCR}/scan.png', 'wb').write(scan_page()); open(f'{SCR}/playa.jpg', 'wb').write(playa)

# HTML para el PDF (lo imprime Chromium)
b64 = lambda p, mime: f'data:{mime};base64,' + base64.b64encode(open(p, 'rb').read()).decode()
html = f'''<!doctype html><html><head><meta charset="utf-8"><style>
body{{font-family:sans-serif}} .pg{{page-break-after:always}} img.f{{width:320px}} img.l{{width:24px}}</style></head><body>
<div class="pg"><img class="l" src="{b64(SCR+'/logo.png','image/png')}"> <b>Inmobiliaria Demo</b>
<h2>Dpto. en Miraflores - Vta.</h2><p>Av. Larco 1150, piso 8. 3 dorm., 2 baños, 1 cochera. Área techada 85 m2.
Precio US$ 185,000. Estado: a estrenar. Vista al mar, gimnasio y piscina del edificio. Cel 987654321</p>
<img class="f" src="{b64(SCR+'/sala.jpg','image/jpeg')}"> <img class="f" src="{b64(SCR+'/cocina.jpg','image/jpeg')}"></div>
<div class="pg"><h2>Casa en La Molina - Alq.</h2><p>Calle Los Pinos 120. Terreno 300 m2, techada 250 m2, 2 pisos. 4 dorm.
S/ 6,500 mensual. E/C (entrega en marzo).</p><img class="f" src="{b64(SCR+'/fachada.jpg','image/jpeg')}"></div>
<div><img style="width:700px" src="{b64(SCR+'/scan.png','image/png')}"></div>
</body></html>'''
open(f'{SCR}/doc.html', 'w').write(html)

# DOCX
doc = Document()
doc.add_heading('Inventario Inmobiliaria Demo', 1)
doc.add_paragraph().add_run().add_picture(f'{SCR}/logo.png', width=Inches(0.3))
doc.add_heading('Depa en San Isidro', 2)
doc.add_paragraph('Calle Las Begonias 450. Venta. 2 dorm, 2 baños. 95 m2 techados. Precio US$ 260 mil. Ascensor: sí.')
t = doc.add_table(rows=2, cols=3)
for i, (a, b) in enumerate([('Estacionamientos', '1'), ('Antigüedad', '2015')]):
    t.cell(i, 0).text = a; t.cell(i, 1).text = b
doc.add_paragraph().add_run().add_picture(f'{SCR}/sala.jpg', width=Inches(3))
doc.add_page_break()
doc.add_heading('Local comercial en Surquillo', 2)
doc.add_paragraph('Av. Angamos 1200. Alquiler. 60 m2. S/. 4,800. Preventa no aplica.')
doc.add_paragraph().add_run().add_picture(f'{SCR}/fachada.jpg', width=Inches(3))
doc.save(f'{OUT}/inventario.docx')

# PPTX
prs = Presentation()
blank = prs.slide_layouts[6]
s1 = prs.slides.add_slide(blank)
tb = s1.shapes.add_textbox(PInches(0.5), PInches(0.3), PInches(9), PInches(1.5)).text_frame
tb.text = 'Casa de playa en Punta Hermosa'
p = tb.add_paragraph(); p.text = 'Venta · US$ 420,000 · Terreno 400 m2 · Techada 280 m2 · 5 dorm · 4 baños · Pet friendly'
s1.shapes.add_picture(f'{SCR}/playa.jpg', PInches(0.5), PInches(2), width=PInches(4))
s1.shapes.add_picture(f'{SCR}/logo.png', PInches(9), PInches(6.8), width=PInches(0.4))
s1.notes_slide.notes_text_frame.text = 'Nota: incluye 2 estacionamientos.'
s2 = prs.slides.add_slide(blank)
tb2 = s2.shapes.add_textbox(PInches(0.5), PInches(0.3), PInches(9), PInches(1.5)).text_frame
tb2.text = 'Of. en San Isidro - Alq.'
p = tb2.add_paragraph(); p.text = 'Av. Camino Real 390, piso 12 · 120 m2 · US$ 3,200 · En const.'
s2.shapes.add_picture(f'{SCR}/cocina.jpg', PInches(0.5), PInches(2), width=PInches(4))
prs.save(f'{OUT}/inventario.pptx')
print('ok')
